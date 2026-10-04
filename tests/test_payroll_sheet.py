# tests/test_payroll_sheet.py
import pytest
import os
import openpyxl
from report_service import ReportService

class TestPayrollSheetAndExcel:
    """اختبارات كشف المسير المجمع وتصدير Excel"""

    def test_generate_monthly_payroll_sheet_pdf(self, sample_crew_data, tmp_path):
        out_pdf = str(tmp_path / "test_payroll_sheet.pdf")
        crew_list = [sample_crew_data]

        res_path = ReportService.generate_monthly_payroll_sheet(crew_list, 9, 2026, output_path=out_pdf)

        assert os.path.exists(res_path)
        assert os.path.getsize(res_path) > 0

    def test_export_payroll_to_excel(self, sample_crew_data, tmp_path):
        out_xlsx = str(tmp_path / "test_payroll.xlsx")
        crew_list = [sample_crew_data]

        res_path = ReportService.export_payroll_to_excel(crew_list, 9, 2026, output_path=out_xlsx)

        assert os.path.exists(res_path)
        assert os.path.getsize(res_path) > 0

        wb = openpyxl.load_workbook(res_path)
        assert f"Payroll_2026_09" in wb.sheetnames
        ws = wb[f"Payroll_2026_09"]
        assert ws["B4"].value == sample_crew_data["name"]

    def test_export_master_cash_to_excel(self, tmp_path):
        out_xlsx = str(tmp_path / "test_master_cash.xlsx")
        txs = [
            {"date": "2026-09-01", "description": "عهدة من الوكيل", "type": "وارد", "amount": 5000.0},
            {"date": "2026-09-03", "description": "مشتريات طعام", "type": "صادر", "amount": 800.0}
        ]
        summary = {"in": 5000.0, "out": 800.0, "crew_adv": 500.0, "net": 3700.0}

        res_path = ReportService.export_master_cash_to_excel(txs, summary, 9, 2026, output_path=out_xlsx)

        assert os.path.exists(res_path)
        assert os.path.getsize(res_path) > 0

        wb = openpyxl.load_workbook(res_path)
        assert "MasterCash_2026_09" in wb.sheetnames
        ws = wb["MasterCash_2026_09"]
        assert ws["B3"].value == 5000.0
