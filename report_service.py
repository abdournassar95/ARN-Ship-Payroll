# ... existing code ...
from PyQt6.QtWidgets import QApplication
from PyQt6.QtGui import QTextDocument, QPageSize, QPageLayout
from PyQt6.QtPrintSupport import QPrinter
from PyQt6.QtCore import QSizeF, QMarginsF
import html
import tempfile
import os
from datetime import datetime
from typing import List, Dict, Any, Optional

class ReportService:
    _app = None

    @classmethod
    def _ensure_app(cls):
        if cls._app is None:
            if QApplication.instance():
                cls._app = QApplication.instance()
            else:
                import sys
                cls._app = QApplication(sys.argv)

    @classmethod
    def _normalize_crew_data(cls, crew: Dict[str, Any]) -> Dict[str, Any]:
        if not isinstance(crew, dict):
            crew = {}
            
        ui_total_due = float(crew.get('total_due', 0))
        curr_extra = float(crew.get('curr_extra', crew.get('extra', 0)))
        curr_ded = float(crew.get('curr_ded', crew.get('deduction', 0)))
        prev_balance = float(crew.get('prev_balance', 0))
        
        # Fetch company and vessel name if not present
        company_name = crew.get('company_name')
        vessel_name = crew.get('vessel_name')
        if not company_name or not vessel_name:
            import sqlite3
            try:
                conn = sqlite3.connect('arn_ship_payroll.db')
                c = conn.cursor()
                c.execute("SELECT company_name, vessel_name FROM system_settings LIMIT 1")
                row = c.fetchone()
                if row:
                    company_name = company_name or row[0] or "ARN FLEET MANAGEMENT"
                    vessel_name = vessel_name or row[1] or "Vessel"
                conn.close()
            except Exception:
                pass
        
        company_name = company_name or "ARN FLEET MANAGEMENT"
        vessel_name = vessel_name or "Vessel"
        
        # إجمالي المستحق القادم من محرك الرواتب (PayrollEngine) يساوي:
        #     أجر الفترة + إضافي + رصيد سابق − خصم مباشر
        # لذلك نستخرج «أجر الفترة» ثم نُعيد تركيب الإجمالي بنفس المعادلة،
        # وبذلك يطابق الكشف محرك الرواتب حرفياً (بدل إعادة إدخال الخصم في الإجمالي).
        period_wage = ui_total_due - curr_extra - prev_balance + curr_ded
        total_earnings = period_wage + curr_extra + prev_balance - curr_ded
        
        return {
            'company_name': str(company_name),
            'vessel_name': str(vessel_name),
            'name': str(crew.get('name', 'غير محدد')),
            'rank': str(crew.get('rank', 'غير محدد')),
            'month': int(crew.get('month', datetime.now().month)),
            'year': int(crew.get('year', datetime.now().year)),
            'basic_wage': float(crew.get('basic_wage', crew.get('wage', crew.get('MonthlyWage', crew.get('monthly_wage', 0))))),
            'worked_days': float(crew.get('worked_days', crew.get('net_days', 0))),
            'total_due': period_wage,
            'extra': curr_extra,
            'prev_balance': prev_balance,
            'total_earnings': total_earnings,
            'deduction': curr_ded,
            'payment_cash': float(crew.get('payment_cash', crew.get('cum_cash', 0))),
            'cigarette': float(crew.get('cigarette', crew.get('cum_cig', 0))),
            'transfer': float(crew.get('transfer', crew.get('cum_trans', 0))),
            'total_received': float(crew.get('total_received', 0)),
            'final_balance': float(crew.get('final_balance', 0))
        }

    @staticmethod
    def _format_currency(value: float) -> str:
        """تنسيق العملة مع فاصلة الآلاف."""
        return f"{value:,.2f}"

    @classmethod
    def _common_css(cls) -> str:
        """
        تنسيقات CSS مصممة خصيصاً لتجاوز قيود QTextDocument في Qt.
        تم الاعتماد أكثر على خصائص HTML المضمنة (Inline) لضمان دقة الطباعة.
        """
        return """
        body {
            font-family: 'Cairo', 'Segoe UI', Tahoma, Arial, sans-serif;
            direction: rtl;
            margin: 0;
            padding: 0;
            color: #0F172A;
        }
        .title-text {
            font-size: 22pt;
            font-weight: bold;
            color: #1E3A8A;
        }
        .subtitle-text {
            font-size: 14pt;
            font-weight: bold;
            color: #475569;
        }
        .box-outline {
            border: 2px solid #1E3A8A;
            border-radius: 8px;
            padding: 15px;
            background-color: #ffffff;
        }
        .data-label {
            font-size: 11pt;
            color: #334155;
        }
        .data-val {
            font-size: 11pt;
            font-weight: bold;
            color: #0F172A;
        }
        .kpi-title {
            font-size: 11pt;
            color: #475569;
            font-weight: bold;
        }
        .kpi-val {
            font-size: 16pt;
            font-weight: bold;
            color: #0F172A;
        }
        .footer-text {
            font-size: 10pt;
            color: #64748B;
        }
        """

    @classmethod
    def _build_payslip_html(cls, crew: Dict[str, Any]) -> str:
        c = cls._normalize_crew_data(crew)
        
        month_names = {1: 'يناير', 2: 'فبراير', 3: 'مارس', 4: 'أبريل', 5: 'مايو', 6: 'يونيو', 
                       7: 'يوليو', 8: 'أغسطس', 9: 'سبتمبر', 10: 'أكتوبر', 11: 'نوفمبر', 12: 'ديسمبر'}
        month_str = month_names.get(c['month'], str(c['month']))
        
        # حماية من أي نصوص قد تكسر الـ HTML
        safe_name = html.escape(c['name'])
        safe_rank = html.escape(c['rank'])
        safe_company = html.escape(c['company_name'])
        safe_vessel = html.escape(c['vessel_name'])

        # سطر الرصيد السابق يظهر فقط عند وجود رصيد مرحّل (موجب أو سالب)
        prev_row_html = ""
        if abs(c['prev_balance']) > 0.004:
            prev_color = "#059669" if c['prev_balance'] >= 0 else "#DC2626"
            prev_value = (
                f'({cls._format_currency(abs(c["prev_balance"]))})'
                if c['prev_balance'] < 0
                else cls._format_currency(c['prev_balance'])
            )
            prev_row_html = (
                '<tr>'
                '<td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">'
                'رصيد سابق (Previous Balance)</td>'
                f'<td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold; color: {prev_color};">'
                f'${prev_value}</td>'
                '</tr>'
            )

        return f"""
        <div style="padding: 10px; font-family: 'Cairo', sans-serif;">
            <!-- Header Card -->
            <table width="100%" cellpadding="8" cellspacing="0" style="border: 1px solid #94A3B8; border-radius: 6px; margin-bottom: 15px;">
                <tr>
                    <td align="right" width="50%" style="text-align: right;">
                        <div style="font-size: 18pt; font-weight: bold; color: #1E293B;">{safe_company}</div>
                        <div style="font-size: 14pt; font-weight: bold; color: #334155; margin-top: 5px;">M/V {safe_vessel}</div>
                    </td>
                    <td align="left" width="50%" valign="top" style="text-align: left;">
                        <div style="font-size: 16pt; font-weight: bold; color: #1E293B;">إيصال رواتب (Payslip)</div>
                        <div style="font-size: 12pt; color: #475569; margin-top: 5px;">{month_str} {c['year']}</div>
                    </td>
                </tr>
            </table>

            <!-- Crew Info Card -->
            <table width="100%" cellpadding="8" cellspacing="0" style="border: 1px solid #94A3B8; border-radius: 6px; margin-bottom: 20px;">
                <tr>
                    <td align="center" style="text-align: center;">
                        <div style="font-size: 16pt; font-weight: bold; color: #0F172A;">{safe_name}</div>
                        <div style="font-size: 12pt; color: #475569; margin-top: 3px; font-weight: bold;">{safe_rank}</div>
                    </td>
                </tr>
            </table>

            <!-- Financials (Earnings & Deductions side by side cards) -->
            <table width="100%" cellspacing="15" cellpadding="0" style="margin-bottom: 15px;">
                <tr>
                    <td width="50%" valign="top">
                        <!-- Earnings Card -->
                        <table width="100%" cellpadding="6" cellspacing="0" style="border: 1px solid #94A3B8; border-radius: 6px;">
                            <tr>
                                <td colspan="2" align="center" style="border-bottom: 1px solid #94A3B8; font-size: 12pt; font-weight: bold; color: #1E293B; padding-bottom: 8px;">المستحقات (Earnings)</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">الراتب الأساسي (Basic Salary)</td>
                                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold;">${cls._format_currency(c['basic_wage'])}</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">أيام العمل (Net Days)</td>
                                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold;">{c['worked_days']}</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">أجر الفترة (Period Wage)</td>
                                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold;">${cls._format_currency(c['total_due'])}</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #94A3B8; padding-top: 8px; color: #334155;">مكافآت (Bonus)</td>
                                <td align="left" style="border-bottom: 1px solid #94A3B8; padding-top: 8px; font-weight: bold;">+${cls._format_currency(c['extra'])}</td>
                            </tr>
                            {prev_row_html}
                            <tr>
                                <td align="right" style="padding-top: 10px; font-weight: bold; color: #0F172A; font-size: 11pt;">إجمالي المستحق (Total Due)</td>
                                <td align="left" style="padding-top: 10px; font-weight: bold; color: #0F172A; font-size: 11pt;">${cls._format_currency(c['total_earnings'])}</td>
                            </tr>
                        </table>
                    </td>
                    <td width="50%" valign="top">
                        <!-- Deductions Card -->
                        <table width="100%" cellpadding="6" cellspacing="0" style="border: 1px solid #94A3B8; border-radius: 6px;">
                            <tr>
                                <td colspan="2" align="center" style="border-bottom: 1px solid #94A3B8; font-size: 12pt; font-weight: bold; color: #1E293B; padding-bottom: 8px;">الخصومات والمستلم (Deductions)</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">خصم مباشر (Deduction)</td>
                                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold;">(${cls._format_currency(c['deduction'])})</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">سلفة نقدية (Cash Advance)</td>
                                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold;">(${cls._format_currency(c['payment_cash'])})</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; color: #334155;">سجائر (Cigarettes)</td>
                                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding-top: 8px; font-weight: bold;">(${cls._format_currency(c['cigarette'])})</td>
                            </tr>
                            <tr>
                                <td align="right" style="border-bottom: 1px solid #94A3B8; padding-top: 8px; color: #334155;">تحويل (Transfer)</td>
                                <td align="left" style="border-bottom: 1px solid #94A3B8; padding-top: 8px; font-weight: bold;">(${cls._format_currency(c['transfer'])})</td>
                            </tr>
                            <tr>
                                <td align="right" style="padding-top: 10px; font-weight: bold; color: #0F172A; font-size: 11pt;">إجمالي المستلم (Received)</td>
                                <td align="left" style="padding-top: 10px; font-weight: bold; color: #0F172A; font-size: 11pt;">(${cls._format_currency(c['total_received'])})</td>
                            </tr>
                        </table>
                    </td>
                </tr>
            </table>

            <!-- Net Balance Card -->
            <table width="100%" cellpadding="12" cellspacing="0" style="border: 2px solid #0F172A; border-radius: 8px; margin-bottom: 25px;">
                <tr>
                    <td align="right" width="50%">
                        <div style="font-size: 16pt; font-weight: bold; color: #0F172A;">صافي الرصيد المستحق</div>
                        <div style="font-size: 12pt; color: #334155; margin-top: 3px;">Net Balance Due</div>
                    </td>
                    <td align="left" width="50%">
                        <div style="font-size: 22pt; font-weight: bold; color: #0F172A; text-align: left;">${cls._format_currency(c['final_balance'])}</div>
                    </td>
                </tr>
            </table>

            <!-- Signatures Card -->
            <table width="100%" cellspacing="20" cellpadding="0" style="margin-top: 70px;">
                <tr>
                    <td width="50%" align="center">
                        <div style="border-top: 1px solid #94A3B8; padding-top: 8px; font-weight: bold; color: #334155; font-size: 11pt; width: 80%;">توقيع البحار (Seaman's Signature)</div>
                    </td>
                    <td width="50%" align="center">
                        <div style="border-top: 1px solid #94A3B8; padding-top: 8px; font-weight: bold; color: #334155; font-size: 11pt; width: 80%;">يعتمد من ربان السفينة (Master's Approval)</div>
                    </td>
                </tr>
            </table>

            <!-- Footer -->
            <table width="100%" cellspacing="0" cellpadding="5" style="margin-top: 40px; border-top: 1px solid #E2E8F0; padding-top: 8px;">
                <tr>
                    <td align="right" style="font-size: 10pt; color: #64748B;">تاريخ الإصدار: {datetime.now().strftime('%Y-%m-%d %H:%M')}</td>
                </tr>
                <tr>
                    <td align="center" style="padding-top: 15px;">
                        <span style="color: #64748B; font-size: 8pt; font-weight: normal;">Powered by</span><br/>
                        <span style="font-size: 11pt; font-weight: bold; color: #0F172A;">ARN Marine Technology</span>
                    </td>
                </tr>
            </table>
        </div>
        """

    @classmethod
    def generate_table_report(cls, title: str, subtitle: str, headers: List[str], rows: List[List[Any]],
                              output_path: str, col_pct: Optional[List[float]] = None,
                              orientation: QPageLayout.Orientation = QPageLayout.Orientation.Landscape,
                              footer_note: str = "ARN Technology — تقرير رسمي مُولَّد آلياً",
                              page_size: QPageSize.PageSizeId = QPageSize.PageSizeId.A4) -> str:
        """
        توليد تقرير جدولي رسمي (HTML ← PDF عبر Qt) بدعم كامل للعربية (RTL).

        يُستخدم في سجل التدقيق ومركز التنبيهات، ويضمن ظهور النص العربي سليماً
        دون الاعتماد على خطوط لاتينية (مثل Helvetica في reportlab) التي تُظهر العربية
        كصناديق فارغة. كما يوحّد شكل التقارير مع كشوف الرواتب القائمة.
        """
        headers = [str(h) for h in headers]
        n = max(1, len(headers))
        if not col_pct or len(col_pct) != n:
            col_pct = [round(100.0 / n, 2)] * n

        th_html = "".join(
            f'<th width="{c}%" style="border: 1px solid #334155; color: #F8FAFC; padding: 6px; font-size: 8.5pt;">{html.escape(h)}</th>'
            for h, c in zip(headers, col_pct)
        )

        tr_html = []
        for i, row in enumerate(rows):
            bg = "#F8FAFC" if i % 2 else "#FFFFFF"
            cells = "".join(
                '<td align="center" style="border: 1px solid #CBD5E1; padding: 4px; font-size: 8pt;">'
                f'{html.escape("" if v is None else str(v))}</td>'
                for v in row
            )
            tr_html.append(f'<tr bgcolor="{bg}">{cells}</tr>')

        if not tr_html:
            tr_html.append(
                f'<tr><td colspan="{n}" align="center" style="padding: 18px; color: #64748B;">لا توجد بيانات للعرض.</td></tr>'
            )

        html_content = f"""<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head><meta charset="UTF-8"></head>
<body style="font-family: 'Cairo', sans-serif; direction: rtl; margin: 0; padding: 5px;">
    <table width="100%" cellpadding="4" cellspacing="0" style="border-bottom: 2px solid #1E3A8A; margin-bottom: 10px;">
        <tr>
            <td align="center">
                <div style="font-size: 16pt; font-weight: bold; color: #1E3A8A;">{html.escape(title)}</div>
                <div style="font-size: 9.5pt; color: #475569; margin-top: 4px;">{html.escape(subtitle)}</div>
            </td>
        </tr>
    </table>
    <table width="100%" cellpadding="3" cellspacing="0" style="border-collapse: collapse; border: 1px solid #CBD5E1;">
        <thead><tr bgcolor="#0F172A">{th_html}</tr></thead>
        <tbody>{''.join(tr_html)}</tbody>
    </table>
    <table width="100%" cellspacing="0" cellpadding="4" style="margin-top: 14px; border-top: 1px solid #E2E8F0;">
        <tr>
            <td align="right" style="font-size: 8.5pt; color: #64748B;">{html.escape(footer_note)}</td>
            <td align="left" style="font-size: 8.5pt; color: #64748B;">تاريخ الإصدار: {datetime.now().strftime('%Y-%m-%d %H:%M')}</td>
        </tr>
    </table>
</body>
</html>"""

        return cls._render_pdf(html_content, output_path, page_size=page_size, orientation=orientation)

    @classmethod
    def _render_pdf(cls, html_content: str, output_path: str, page_size: QPageSize.PageSizeId = QPageSize.PageSizeId.A4,
                    orientation: QPageLayout.Orientation = QPageLayout.Orientation.Portrait) -> str:
        """يقوم بتوليد ملف الـ PDF مع ضبط دقيق للهوامش ودقة الطابعة."""
        cls._ensure_app()

        output_dir = os.path.dirname(output_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir, exist_ok=True)

        doc = QTextDocument()
        doc.setDefaultStyleSheet(cls._common_css())
        doc.setHtml(html_content)

        # استخدام دقة الشاشة العادية لمنع صغر الخطوط جداً عند التصدير للـ PDF
        printer = QPrinter(QPrinter.PrinterMode.ScreenResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(output_path)
        
        # إعداد مقاس الصفحة وتحديد الهوامش والاتجاه (طولي أو عرضي)
        layout = QPageLayout(
            QPageSize(page_size),
            orientation,
            QMarginsF(10, 10, 10, 10),
            QPageLayout.Unit.Millimeter
        )
        printer.setPageLayout(layout)

        rect_pixels = printer.pageLayout().paintRectPixels(printer.resolution())
        doc.setPageSize(QSizeF(float(rect_pixels.width()), float(rect_pixels.height())))

        doc.print(printer)

        if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
            raise RuntimeError(f"فشل إنشاء ملف PDF في المسار: {output_path}")

        return output_path

    @classmethod
    def generate_crew_payslip(cls, crew_data: Dict[str, Any], output_path: Optional[str] = None,
                              page_size: QPageSize.PageSizeId = QPageSize.PageSizeId.A4) -> str:
        data = cls._normalize_crew_data(crew_data)
        
        html_body = cls._build_payslip_html(data)
        html_content = f"""<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head>
    <meta charset="UTF-8">
</head>
<body>
    {html_body}
</body>
</html>"""

        if not output_path:
            import time
            safe_name = "".join(ch for ch in data['name'] if ch.isalnum() or ch in (' ', '_')).strip() or "Crew"
            temp_dir = tempfile.gettempdir()
            ts = int(time.time())
            output_path = os.path.join(temp_dir, f"Payslip_{safe_name}_{data['month']}_{data['year']}_{ts}.pdf")

        return cls._render_pdf(html_content, output_path, page_size)

    @classmethod
    def generate_batch_crew_payslips(cls, crew_data_list: List[Dict[str, Any]],
                                     output_path: Optional[str] = None,
                                     page_size: QPageSize.PageSizeId = QPageSize.PageSizeId.A4) -> str:
        if not crew_data_list:
            raise ValueError("لا توجد بيانات بحارة لطباعتها.")

        pages_html = []
        for crew in crew_data_list:
            data = cls._normalize_crew_data(crew)
            pages_html.append(cls._build_payslip_html(data))

        full_body = '<div style="page-break-before: always;"></div>'.join(pages_html)
        html_content = f"""<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head>
    <meta charset="UTF-8">
</head>
<body>
    {full_body}
</body>
</html>"""

        if not output_path:
            time_stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            temp_dir = tempfile.gettempdir()
            output_path = os.path.join(temp_dir, f"All_Crew_Payslips_{time_stamp}.pdf")

        return cls._render_pdf(html_content, output_path, page_size)

    @classmethod
    def generate_master_cash_report(cls, month: Any, year: Any, transactions: List[Any],
                                    summary: Dict[str, Any], output_path: Optional[str] = None,
                                    page_size: QPageSize.PageSizeId = QPageSize.PageSizeId.A4) -> str:
        """توليد تقرير صندوق القبطان بتصميم متوافق ومحسّن للطباعة عبر Qt."""
        net_val = float(summary.get('net', 0))
        in_val = float(summary.get('in', 0))
        out_val = float(summary.get('out', 0))
        old_adv = float(summary.get('old_adv', 0))

        rows_html = ""
        running_balance = old_adv
        
        for tx in transactions:
            if isinstance(tx, dict):
                date_str = str(tx.get('date', ''))
                desc = html.escape(str(tx.get('desc', tx.get('description', ''))))
                t_type = html.escape(str(tx.get('type', tx.get('transaction_type', 'وارد'))))
                amount = float(tx.get('amount', 0))
            else:
                date_str = str(tx[0]) if len(tx) > 0 else ''
                desc = html.escape(str(tx[1]) if len(tx) > 1 else '')
                t_type = html.escape(str(tx[2]) if len(tx) > 2 else 'وارد')
                amount = float(tx[3]) if len(tx) > 3 else 0.0

            if t_type == "وارد":
                color = "#059669"
                sign = "+"
                running_balance += amount
            else:
                color = "#DC2626"
                sign = "-"
                running_balance -= amount

            rows_html += f"""
            <tr>
                <td align="center" style="border-bottom: 1px solid #E2E8F0; padding: 8px;">{date_str}</td>
                <td align="right" style="border-bottom: 1px solid #E2E8F0; padding: 8px;">{desc}</td>
                <td align="center" style="border-bottom: 1px solid #E2E8F0; padding: 8px; color:{color}; font-weight:bold;">{t_type}</td>
                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding: 8px; color:{color}; font-weight:bold;">{sign}${cls._format_currency(amount)}</td>
                <td align="left" style="border-bottom: 1px solid #E2E8F0; padding: 8px; font-weight:bold; color: #0F172A;">${cls._format_currency(running_balance)}</td>
            </tr>
            """

        if not rows_html:
            rows_html = '<tr><td colspan="5" align="center" style="padding: 20px; color: #64748B;">لا توجد حركات مالية مسجلة لهذه الفترة.</td></tr>'

        html_content = f"""<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head><meta charset="UTF-8"></head>
<body>
    <div class="box-outline">
        <!-- الهيدر -->
        <table width="100%" cellpadding="5" cellspacing="0" style="border-bottom: 2px solid #1E3A8A; margin-bottom: 15px;">
            <tr>
                <td align="center">
                    <div class="title-text">ARN FLEET MANAGEMENT</div>
                    <div class="subtitle-text">Master Cash Statement / تقرير صندوق القبطان</div>
                    <div style="font-size: 12pt; color: #475569; font-weight: bold; margin-top: 5px;">فترة المحاسبة: {month} / {year}</div>
                </td>
            </tr>
        </table>

        <!-- كروت الإحصائيات -->
        <table width="100%" cellspacing="12" cellpadding="0" style="margin-bottom: 20px;">
            <tr>
                <td width="25%" align="center" bgcolor="#F8FAFC" style="border: 1px solid #CBD5E1; padding: 12px; border-radius: 6px;">
                    <div class="kpi-title">عهدة سابقة</div>
                    <div class="kpi-val">${cls._format_currency(old_adv)}</div>
                </td>
                <td width="25%" align="center" bgcolor="#F8FAFC" style="border: 1px solid #CBD5E1; padding: 12px; border-radius: 6px;">
                    <div class="kpi-title">إجمالي الوارد</div>
                    <div class="kpi-val" style="color: #059669;">+${cls._format_currency(in_val)}</div>
                </td>
                <td width="25%" align="center" bgcolor="#F8FAFC" style="border: 1px solid #CBD5E1; padding: 12px; border-radius: 6px;">
                    <div class="kpi-title">إجمالي الصادر</div>
                    <div class="kpi-val" style="color: #DC2626;">-${cls._format_currency(out_val)}</div>
                </td>
                <td width="25%" align="center" bgcolor="#EFF6FF" style="border: 2px solid #1E3A8A; padding: 12px; border-radius: 6px;">
                    <div class="kpi-title" style="color: #1E3A8A;">الرصيد الصافي</div>
                    <div class="kpi-val" style="color: #1E3A8A;">${cls._format_currency(net_val)}</div>
                </td>
            </tr>
        </table>

        <div style="font-size: 14pt; font-weight: bold; color: #0F172A; margin-bottom: 10px; border-bottom: 1px solid #CBD5E1; padding-bottom: 5px;">
            سجل الحركات التفصيلي:
        </div>

        <!-- جدول الحركات -->
        <table width="100%" cellspacing="0" cellpadding="0" style="border: 1px solid #CBD5E1;">
            <thead>
                <tr>
                    <th width="15%" bgcolor="#F1F5F9" style="padding: 10px; border-bottom: 2px solid #CBD5E1; color: #0F172A;">التاريخ</th>
                    <th width="40%" bgcolor="#F1F5F9" style="padding: 10px; border-bottom: 2px solid #CBD5E1; color: #0F172A;">البيان</th>
                    <th width="15%" bgcolor="#F1F5F9" style="padding: 10px; border-bottom: 2px solid #CBD5E1; color: #0F172A;">النوع</th>
                    <th width="15%" bgcolor="#F1F5F9" style="padding: 10px; border-bottom: 2px solid #CBD5E1; color: #0F172A;">المبلغ</th>
                    <th width="15%" bgcolor="#F1F5F9" style="padding: 10px; border-bottom: 2px solid #CBD5E1; color: #0F172A;">الرصيد</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>

        <!-- التوقيعات -->
        <table width="100%" cellspacing="15" cellpadding="0" style="margin-top: 20px;">
            <tr>
                <td width="50%" align="center" bgcolor="#F8FAFC" style="border: 1px solid #CBD5E1; padding: 25px 10px 10px 10px;">
                    <div style="border-top: 1px dashed #94A3B8; padding-top: 5px; font-weight: bold; color: #334155;">إعداد المحاسب / Prepared By</div>
                </td>
                <td width="50%" align="center" bgcolor="#F8FAFC" style="border: 1px solid #CBD5E1; padding: 25px 10px 10px 10px;">
                    <div style="border-top: 1px dashed #94A3B8; padding-top: 5px; font-weight: bold; color: #334155;">اعتماد الربان / Master's Approval</div>
                </td>
            </tr>
        </table>
    </div>
</body>
</html>"""

        if not output_path:
            temp_dir = tempfile.gettempdir()
            output_path = os.path.join(temp_dir, f"MasterCash_{month}_{year}.pdf")

        return cls._render_pdf(html_content, output_path, page_size)

    @classmethod
    def generate_monthly_payroll_sheet(cls, crew_data_list: List[Dict[str, Any]], month: int, year: int,
                                       output_path: Optional[str] = None) -> str:
        """
        توليد كشف مسير الرواتب الشهري المجمع لكامل الطاقم بصيغة Landscape A4 PDF
        المعتمد رسمياً لهيئات التفتيش البحري والشركات الملاحية
        """
        if not crew_data_list:
            raise ValueError("لا توجد بيانات بحارة لطباعتها في كشف المسير.")

        company_name = "ARN FLEET MANAGEMENT"
        vessel_name = "Vessel"
        import sqlite3
        try:
            with sqlite3.connect('arn_ship_payroll.db') as conn:
                row = conn.execute("SELECT company_name, vessel_name FROM system_settings LIMIT 1").fetchone()
                if row:
                    company_name = row[0] or company_name
                    vessel_name = row[1] or vessel_name
        except Exception:
            pass

        month_names = {1: 'يناير', 2: 'فبراير', 3: 'مارس', 4: 'أبريل', 5: 'مايو', 6: 'يونيو', 
                       7: 'يوليو', 8: 'أغسطس', 9: 'سبتمبر', 10: 'أكتوبر', 11: 'نوفمبر', 12: 'ديسمبر'}
        month_str = month_names.get(month, str(month))

        rows_html = []
        tot_wage = 0.0
        tot_extra = 0.0
        tot_prev = 0.0
        tot_ded = 0.0
        tot_due = 0.0
        tot_cash = 0.0
        tot_cig = 0.0
        tot_trans = 0.0
        tot_recv = 0.0
        tot_net = 0.0

        for i, raw_c in enumerate(crew_data_list):
            c = cls._normalize_crew_data(raw_c)
            b_wage = c['basic_wage']
            days = c['worked_days']
            p_wage = c['total_due']
            extra = c['extra']
            prev = c['prev_balance']
            ded = c['deduction']
            t_due = c['total_earnings']
            c_cash = c['payment_cash']
            c_cig = c['cigarette']
            c_trans = c['transfer']
            t_recv = c['total_received']
            f_bal = c['final_balance']

            tot_wage += b_wage
            tot_extra += extra
            tot_prev += prev
            tot_ded += ded
            tot_due += t_due
            tot_cash += c_cash
            tot_cig += c_cig
            tot_trans += c_trans
            tot_recv += t_recv
            tot_net += f_bal

            bg_row = "#F8FAFC" if i % 2 == 1 else "#FFFFFF"
            bal_color = "#059669" if f_bal >= 0 else "#DC2626"

            rows_html.append(f"""
            <tr bgcolor="{bg_row}">
                <td align="center" style="border: 1px solid #CBD5E1; padding: 4px; font-size: 8pt;">{i + 1}</td>
                <td align="right" style="border: 1px solid #CBD5E1; padding: 4px; font-size: 8pt; font-weight: bold;">{html.escape(c['name'])}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 4px; font-size: 8pt;">{html.escape(c['rank'])}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; white-space: nowrap;">${cls._format_currency(b_wage)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 4px; font-size: 8pt;">{days}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; white-space: nowrap;">${cls._format_currency(p_wage)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; color: #059669; white-space: nowrap;">${cls._format_currency(extra)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; color: {'#059669' if prev >= 0 else '#DC2626'}; white-space: nowrap;">${'(' + cls._format_currency(abs(prev)) + ')' if prev < 0 else cls._format_currency(prev)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; color: #DC2626; white-space: nowrap;">${cls._format_currency(ded)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; font-weight: bold; color: #1E3A8A; white-space: nowrap;">${cls._format_currency(t_due)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; white-space: nowrap;">${cls._format_currency(c_cash)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; white-space: nowrap;">${cls._format_currency(c_cig)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; white-space: nowrap;">${cls._format_currency(c_trans)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; white-space: nowrap;">${cls._format_currency(t_recv)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 3px 2px; font-size: 7.5pt; font-weight: bold; color: {bal_color}; white-space: nowrap;">${cls._format_currency(f_bal)}</td>
                <td align="center" style="border: 1px solid #CBD5E1; padding: 4px; font-size: 7pt; color: #94A3B8;">&nbsp;</td>
            </tr>
            """)

        table_rows_str = "\n".join(rows_html)

        html_content = f"""<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head>
    <meta charset="UTF-8">
</head>
<body style="font-family: 'Cairo', sans-serif; direction: rtl; margin: 0; padding: 5px;">
    <!-- الهيدر الرسمي -->
    <table width="100%" cellpadding="4" cellspacing="0" style="border-bottom: 2px solid #1E3A8A; margin-bottom: 8px;">
        <tr>
            <td width="35%" align="right">
                <div style="font-size: 15pt; font-weight: bold; color: #1E293B;">{html.escape(company_name)}</div>
                <div style="font-size: 12pt; color: #475569; font-weight: bold;">M/V {html.escape(vessel_name)}</div>
            </td>
            <td width="30%" align="center">
                <div style="font-size: 16pt; font-weight: bold; color: #1E3A8A;">كشف مسير رواتب الطاقم الشهري</div>
                <div style="font-size: 10pt; color: #64748B;">CREW MONTHLY PAYROLL STATEMENT</div>
            </td>
            <td width="35%" align="left" style="text-align: left;">
                <div style="font-size: 12pt; font-weight: bold; color: #1E293B;">فترة الكشف: {month_str} {year}</div>
                <div style="font-size: 9pt; color: #64748B;">تاريخ الإصدار: {datetime.now().strftime('%Y-%m-%d')}</div>
            </td>
        </tr>
    </table>

    <!-- جدول المسير المجمع الشامل -->
    <table width="100%" cellpadding="3" cellspacing="0" style="border-collapse: collapse; border: 1px solid #CBD5E1;">
        <thead>
            <tr bgcolor="#0F172A">
                <th width="3%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">م</th>
                <th width="10%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">اسم البحار</th>
                <th width="6%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">الرتبة</th>
                <th width="6%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">الراتب $</th>
                <th width="4%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">أيام</th>
                <th width="6%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">أجر الفترة</th>
                <th width="5%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">إضافي</th>
                <th width="6%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">رصيد سابق</th>
                <th width="6%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">خصم مباشر</th>
                <th width="7%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">إجمالي المستحق</th>
                <th width="5%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">سلف كاش</th>
                <th width="5%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">سجائر</th>
                <th width="5%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">تحويل</th>
                <th width="7%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">إجمالي المستلم</th>
                <th width="6%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">الصافي المتبقي</th>
                <th width="12%" style="border: 1px solid #334155; color: #F8FAFC; padding: 5px; font-size: 8pt;">توقيع</th>
            </tr>
        </thead>
        <tbody>
            {table_rows_str}
            <!-- صف الإجماليات -->
            <tr bgcolor="#E2E8F0" style="font-weight: bold;">
                <td colspan="3" align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt;">الإجمالي الكلي (Total Summary)</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; white-space: nowrap;">${cls._format_currency(tot_wage)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt;">-</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt;">-</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; color: #059669; white-space: nowrap;">${cls._format_currency(tot_extra)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; white-space: nowrap;">${'(' + cls._format_currency(abs(tot_prev)) + ')' if tot_prev < 0 else cls._format_currency(tot_prev)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; color: #DC2626; white-space: nowrap;">${cls._format_currency(tot_ded)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; color: #1E3A8A; white-space: nowrap;">${cls._format_currency(tot_due)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; white-space: nowrap;">${cls._format_currency(tot_cash)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; white-space: nowrap;">${cls._format_currency(tot_cig)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; white-space: nowrap;">${cls._format_currency(tot_trans)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; white-space: nowrap;">${cls._format_currency(tot_recv)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt; color: #1E3A8A; white-space: nowrap;">${cls._format_currency(tot_net)}</td>
                <td align="center" style="border: 1px solid #94A3B8; padding: 4px 2px; font-size: 8pt;">-</td>
            </tr>
        </tbody>
    </table>

    <!-- صندوق التوقيعات والاعتماد الرسمي -->
    <table width="100%" cellpadding="0" cellspacing="12" style="margin-top: 15px;">
        <tr>
            <td width="33%" align="center" style="border: 1px solid #CBD5E1; border-radius: 6px; padding: 12px;">
                <div style="font-size: 9pt; color: #64748B;">إعداد المحاسب / كاتب الرواتب</div>
                <div style="font-size: 10pt; font-weight: bold; color: #1E293B; margin-top: 3px;">Prepared by Accountant</div>
                <div style="border-top: 1px dashed #94A3B8; margin-top: 25px; padding-top: 4px; font-size: 8pt; color: #64748B;">التوقيع: .......................................</div>
            </td>
            <td width="34%" align="center" style="border: 1px solid #CBD5E1; border-radius: 6px; padding: 12px;">
                <div style="font-size: 9pt; color: #64748B;">مراجعة كبير المهندسين / الضابط الأول</div>
                <div style="font-size: 10pt; font-weight: bold; color: #1E293B; margin-top: 3px;">Checked by Ch.Eng / Ch.Off</div>
                <div style="border-top: 1px dashed #94A3B8; margin-top: 25px; padding-top: 4px; font-size: 8pt; color: #64748B;">التوقيع: .......................................</div>
            </td>
            <td width="33%" align="center" style="border: 1px solid #CBD5E1; border-radius: 6px; padding: 12px;">
                <div style="font-size: 9pt; color: #64748B;">اعتماد ربان السفينة وخاتم السفينة</div>
                <div style="font-size: 10pt; font-weight: bold; color: #1E293B; margin-top: 3px;">Master's Approval & Ship Seal</div>
                <div style="border-top: 1px dashed #94A3B8; margin-top: 25px; padding-top: 4px; font-size: 8pt; color: #64748B;">التوقيع والخاتم: .......................................</div>
            </td>
        </tr>
    </table>
</body>
</html>"""

        if not output_path:
            ts = datetime.now().strftime('%Y%m%d_%H%M%S')
            temp_dir = tempfile.gettempdir()
            output_path = os.path.join(temp_dir, f"Payroll_Sheet_{month}_{year}_{ts}.pdf")

        return cls._render_pdf(
            html_content, output_path,
            page_size=QPageSize.PageSizeId.A4,
            orientation=QPageLayout.Orientation.Landscape
        )

    @classmethod
    def export_payroll_to_excel(cls, crew_data_list: List[Dict[str, Any]], month: int, year: int, output_path: str) -> str:
        """تصدير كشف مسير الرواتب الشهري إلى ملف Excel (.xlsx) منسق باحترافية"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Payroll_{year}_{month:02d}"
        ws.sheet_view.rightToLeft = True

        navy_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Segoe UI", size=10)
        font_bold = Font(name="Segoe UI", size=10, bold=True)
        font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")

        thin_side = Side(style="thin", color="CBD5E1")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.merge_cells("A1:P1")
        ws["A1"] = f"كشف مسير رواتب الطاقم الشهري - {month:02d}/{year}"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        headers = [
            "م", "اسم البحار", "الرتبة", "الراتب الأساسي ($)", "أيام العمل",
            "أجر الفترة ($)", "إضافي ($)", "رصيد سابق ($)", "خصم مباشر ($)",
            "إجمالي المستحق ($)", "سلف كاش ($)", "سجائر ($)", "تحويل ($)",
            "إجمالي المستلم ($)", "صافي الرصيد ($)", "حالة التسوية"
        ]
        money_cols = [4, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15]

        ws.append([]) # row 2
        ws.append(headers) # row 3
        ws.row_dimensions[3].height = 25

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col_num)
            cell.fill = navy_header_fill
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_all

        for i, raw_c in enumerate(crew_data_list):
            c = cls._normalize_crew_data(raw_c)
            row_idx = 4 + i
            is_settled = "مسدد 🔒" if raw_c.get('is_settled') else "قيد العمل 🔓"

            row_vals = [
                i + 1,
                c['name'],
                c['rank'],
                c['basic_wage'],
                c['worked_days'],
                c['total_due'],
                c['extra'],
                c['prev_balance'],
                c['deduction'],
                c['total_earnings'],
                c['payment_cash'],
                c['cigarette'],
                c['transfer'],
                c['total_received'],
                c['final_balance'],
                is_settled
            ]
            ws.append(row_vals)
            ws.row_dimensions[row_idx].height = 20

            for col_num, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = font_data
                cell.border = border_all
                if i % 2 == 1:
                    cell.fill = zebra_fill

                if col_num in money_cols:
                    cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_num == 2:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.font = font_bold
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        sum_row = 4 + len(crew_data_list)
        ws.append(["الإجمالي الكلي", "", ""] + [
            f"=SUM({get_column_letter(col)}4:{get_column_letter(col)}{sum_row-1})"
            if col in money_cols else ""
            for col in range(4, 17)
        ])
        ws.merge_cells(f"A{sum_row}:C{sum_row}")
        ws.row_dimensions[sum_row].height = 24

        for col_num in range(1, 17):
            cell = ws.cell(row=sum_row, column=col_num)
            cell.fill = summary_fill
            cell.font = font_bold
            cell.border = border_all
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_num in money_cols:
                cell.number_format = '$#,##0.00'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        ws.column_dimensions['B'].width = 28

        wb.save(output_path)
        return output_path

    @classmethod
    def export_master_cash_to_excel(cls, transactions: List[Any], summary: Dict[str, Any], month: int, year: int, output_path: str) -> str:
        """تصدير كشف صندوق القبطان إلى ملف Excel (.xlsx) منسق"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"MasterCash_{year}_{month:02d}"
        ws.sheet_view.rightToLeft = True

        navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Segoe UI", size=10)
        font_bold = Font(name="Segoe UI", size=10, bold=True)
        font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")

        thin_side = Side(style="thin", color="CBD5E1")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.merge_cells("A1:E1")
        ws["A1"] = f"تقرير صندوق القبطان (Master Cash) - {month:02d}/{year}"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws["A3"] = "إجمالي الوارد:"
        ws["B3"] = float(summary.get('in', 0))
        ws["B3"].number_format = '$#,##0.00'
        ws["C3"] = "المصروفات:"
        ws["D3"] = float(summary.get('out', 0))
        ws["D3"].number_format = '$#,##0.00'
        ws["A4"] = "سلف البحارة:"
        ws["B4"] = float(summary.get('crew_adv', 0))
        ws["B4"].number_format = '$#,##0.00'
        ws["C4"] = "الرصيد الصافي:"
        ws["D4"] = float(summary.get('net', 0))
        ws["D4"].number_format = '$#,##0.00'
        ws["D4"].font = font_bold

        for r in range(3, 5):
            for c in range(1, 5):
                ws.cell(row=r, column=c).border = border_all
                ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center")

        headers = ["م", "التاريخ", "البيان", "النوع", "المبلغ ($)"]
        ws.append([]) # row 5
        ws.append(headers) # row 6
        ws.row_dimensions[6].height = 24

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=6, column=col_num)
            cell.fill = navy_fill
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border_all

        for i, tx in enumerate(transactions):
            row_idx = 7 + i
            amt = float(tx.get('amount', 0) if isinstance(tx, dict) else (tx[3] if len(tx)>3 else 0))
            t_type = tx.get('type', '') if isinstance(tx, dict) else (tx[2] if len(tx)>2 else '')
            desc = tx.get('description', '') if isinstance(tx, dict) else (tx[1] if len(tx)>1 else '')
            dt = tx.get('date', '') if isinstance(tx, dict) else (tx[0] if len(tx)>0 else '')

            row_vals = [i + 1, dt, desc, t_type, amt]
            ws.append(row_vals)
            ws.row_dimensions[row_idx].height = 20

            for col_num, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = font_data
                cell.border = border_all
                if i % 2 == 1: cell.fill = zebra_fill
                if col_num == 5:
                    cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_num == 3:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        ws.column_dimensions['C'].width = 35

        wb.save(output_path)
        return output_path

    @classmethod
    def generate_documents_report(cls, docs_data: List[Dict[str, Any]], filter_title: str = "كافة أفراد الطاقم", output_path: Optional[str] = None) -> str:
        """
        إنشاء تقرير فحص ومتابعة الشهادات والوثائق البحرية للطاقم (STCW Matrix PDF)
        جاهز للتفتيش الملاحي (Port State Control / Vetting Inspection)
        """
        cls._ensure_app()
        if not output_path:
            out_dir = os.path.join(tempfile.gettempdir(), 'arn_ship_payroll_reports')
            os.makedirs(out_dir, exist_ok=True)
            output_path = os.path.join(out_dir, f"crew_documents_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")

        # بيانات السفينة
        company_name = "ARN FLEET MANAGEMENT"
        vessel_name = "Vessel"
        import sqlite3
        try:
            with sqlite3.connect('arn_ship_payroll.db') as conn:
                c = conn.cursor()
                c.execute("SELECT company_name, vessel_name FROM system_settings LIMIT 1")
                row = c.fetchone()
                if row:
                    company_name = row[0] or company_name
                    vessel_name = row[1] or vessel_name
        except Exception:
            pass

        today = datetime.now().date()
        print_date = datetime.now().strftime('%Y-%m-%d %H:%M')

        # إحصائيات
        total_docs = len(docs_data)
        valid_docs = 0
        expiring_docs = 0
        expired_docs = 0

        rows_html = ""
        for i, d in enumerate(docs_data):
            c_name = html.escape(str(d.get('crew_name', '')))
            c_rank = html.escape(str(d.get('crew_rank', '')))
            d_type = html.escape(str(d.get('doc_type', '')))
            d_num = html.escape(str(d.get('doc_number') or '-'))
            d_iss = html.escape(str(d.get('issue_date') or '-'))
            d_exp = html.escape(str(d.get('expiry_date') or '-'))

            days_left = 999
            try:
                exp_dt = datetime.strptime(str(d.get('expiry_date', ''))[:10], "%Y-%m-%d").date()
                days_left = (exp_dt - today).days
            except Exception:
                pass

            if days_left < 0:
                expired_docs += 1
                status_text = f"منتهية ({abs(days_left)} يوم)"
                badge_style = "background-color: #fee2e2; color: #991b1b; font-weight: bold; border-radius: 4px; padding: 2px 6px;"
            elif days_left <= 60:
                expiring_docs += 1
                status_text = f"تنتهي خلال {days_left} يوم"
                badge_style = "background-color: #fef3c7; color: #92400e; font-weight: bold; border-radius: 4px; padding: 2px 6px;"
            else:
                valid_docs += 1
                status_text = f"سارية ({days_left} يوم)"
                badge_style = "background-color: #dcfce7; color: #166534; font-weight: bold; border-radius: 4px; padding: 2px 6px;"

            bg_color = "#f8fafc" if i % 2 == 1 else "#ffffff"

            rows_html += f"""
            <tr style="background-color: {bg_color}; font-size: 8.5pt;">
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: center;">{i + 1}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: right; font-weight: bold;">{c_name}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: center; color: #0284c7; font-weight: bold;">{c_rank}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: right;">{d_type}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: center;">{d_num}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: center; color: #64748b;">{d_iss}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: center; font-weight: bold;">{d_exp}</td>
                <td style="border: 1px solid #cbd5e1; padding: 5px; text-align: center;"><span style="{badge_style}">{status_text}</span></td>
            </tr>
            """

        html_content = f"""
        <!DOCTYPE html>
        <html dir="rtl" lang="ar">
        <head>
            <meta charset="utf-8">
            <style>
                body {{ font-family: 'Segoe UI', Tahoma, Arial, sans-serif; direction: rtl; margin: 0; padding: 10px; color: #0f172a; }}
                table {{ border-collapse: collapse; width: 100%; }}
                th {{ background-color: #1e3a8a; color: #ffffff; font-weight: bold; font-size: 9pt; padding: 6px 4px; text-align: center; border: 1px solid #1e3a8a; }}
            </style>
        </head>
        <body>
            <table style="margin-bottom: 12px; border: none;">
                <tr>
                    <td style="text-align: right; width: 35%; border: none;">
                        <span style="font-size: 13pt; font-weight: bold; color: #1e3a8a;">{company_name}</span><br>
                        <span style="font-size: 10pt; color: #475569;">سفينة: <b>{vessel_name}</b></span>
                    </td>
                    <td style="text-align: center; width: 35%; border: none;">
                        <span style="font-size: 14pt; font-weight: bold; color: #0f172a;">سجل الشهادات والوثائق الملاحية للطاقم</span><br>
                        <span style="font-size: 9.5pt; color: #64748b;">نظام المتابعة والتوافق مع اتفاقية STCW الدولية</span><br>
                        <span style="font-size: 8.5pt; color: #0284c7; font-weight: bold;">نطاق التقرير: {filter_title}</span>
                    </td>
                    <td style="text-align: left; width: 30%; border: none;">
                        <span style="font-size: 8.5pt; color: #64748b;">تاريخ الطباعة: {print_date}</span><br>
                        <span style="font-size: 8.5pt; color: #10b981; font-weight: bold;">تطبيق ARN Ship Payroll</span>
                    </td>
                </tr>
            </table>

            <!-- إحصائيات سريعة -->
            <table style="margin-bottom: 14px; border: 1px solid #cbd5e1; background-color: #f1f5f9; border-radius: 6px;">
                <tr>
                    <td style="padding: 6px; text-align: center; border: none; font-size: 9.5pt;">
                        إجمالي الوثائق: <b style="color: #0284c7;">{total_docs}</b>
                    </td>
                    <td style="padding: 6px; text-align: center; border: none; font-size: 9.5pt;">
                        شهادات سارية: <b style="color: #166534;">{valid_docs}</b>
                    </td>
                    <td style="padding: 6px; text-align: center; border: none; font-size: 9.5pt;">
                        تنتهي خلال 60 يوماً: <b style="color: #d97706;">{expiring_docs}</b>
                    </td>
                    <td style="padding: 6px; text-align: center; border: none; font-size: 9.5pt;">
                        منتهية الصلاحية: <b style="color: #dc2626;">{expired_docs}</b>
                    </td>
                </tr>
            </table>

            <!-- جدول الشهادات -->
            <table>
                <thead>
                    <tr>
                        <th style="width: 4%;">م</th>
                        <th style="width: 17%;">اسم البحار</th>
                        <th style="width: 10%;">الرتبة</th>
                        <th style="width: 25%;">نوع الشهادة / الوثيقة</th>
                        <th style="width: 12%;">رقم الوثيقة</th>
                        <th style="width: 10%;">تاريخ الإصدار</th>
                        <th style="width: 10%;">تاريخ الانتهاء</th>
                        <th style="width: 12%;">الحالة</th>
                    </tr>
                </thead>
                <tbody>
                    {rows_html}
                </tbody>
            </table>

            <!-- التوقيعات والاعتماد -->
            <table style="margin-top: 30px; border: none;">
                <tr>
                    <td style="text-align: right; width: 35%; border: none; font-size: 10pt;">
                        <b>ضابط السلامة / الإدارة:</b><br><br>
                        ..........................................
                    </td>
                    <td style="text-align: center; width: 30%; border: none; font-size: 10pt;">
                        <b>خاتم السفينة (Ship's Stamp):</b><br><br>
                        [ &nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp; ]
                    </td>
                    <td style="text-align: left; width: 35%; border: none; font-size: 10pt;">
                        <b>ربان السفينة (Master's Signature):</b><br><br>
                        ..........................................
                    </td>
                </tr>
            </table>
        </body>
        </html>
        """

        doc = QTextDocument()
        doc.setHtml(html_content)

        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(output_path)
        printer.setPageOrientation(QPageLayout.Orientation.Landscape)
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        margins = QMarginsF(10, 10, 10, 10)
        printer.setPageMargins(margins, QPageLayout.Unit.Millimeter)

        doc.setPageSize(QSizeF(printer.pageRect(QPrinter.Unit.Point).size()))
        doc.print(printer)

        return output_path
